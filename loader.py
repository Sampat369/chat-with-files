from pathlib import Path
import json
import pandas as pd
from pypdf import PdfReader
from docx import Document
from pptx import Presentation

# Files that can be read directly as text
TEXT_EXTENSIONS = {
    ".txt",
    ".md",
    ".py",
    ".js",
    ".jsx",
    ".ts",
    ".tsx",
    ".java",
    ".c",
    ".h",
    ".cpp",
    ".hpp",
    ".cc",
    ".cs",
    ".go",
    ".rs",
    ".php",
    ".rb",
    ".swift",
    ".kt",
    ".kts",
    ".html",
    ".htm",
    ".css",
    ".scss",
    ".sql",
    ".sh",
    ".bat",
    ".ps1",
    ".yaml",
    ".yml",
    ".toml",
    ".ini",
    ".cfg",
    ".env",
    ".xml",
    ".svg",
    ".log",
}


def load_text_file(path):
    try:
        return Path(path).read_text(
            encoding="utf-8",
            errors="ignore"
        )
    except Exception as e:
        return f"[Could not read file: {e}]"


def load_pdf(path):
    try:
        reader = PdfReader(path)

        text = []

        for page_number, page in enumerate(reader.pages, start=1):
            page_text = page.extract_text() or ""

            text.append(
                f"\n--- Page {page_number} ---\n"
                f"{page_text}"
            )

        return "\n".join(text)

    except Exception as e:
        return f"[Could not read PDF: {e}]"

def load_docx(path):
    try:
        document = Document(path)

        paragraphs = []

        for paragraph in document.paragraphs:
            if paragraph.text.strip():
                paragraphs.append(paragraph.text)

        return "\n".join(paragraphs)

    except Exception as e:
        return f"[Could not read DOCX: {e}]"

def load_docx(path):
    try:
        document = Document(path)

        paragraphs = []

        for paragraph in document.paragraphs:
            if paragraph.text.strip():
                paragraphs.append(paragraph.text)

        return "\n".join(paragraphs)

    except Exception as e:
        return f"[Could not read DOCX: {e}]"
    
def load_csv(path):
    try:
        df = pd.read_csv(path)

        return (
            f"CSV file\n"
            f"Rows: {len(df)}\n"
            f"Columns: {list(df.columns)}\n\n"
            f"{df.to_string()}"
        )

    except Exception as e:
        return f"[Could not read CSV: {e}]"


def load_excel(path):
    try:
        excel = pd.ExcelFile(path)

        output = []

        for sheet in excel.sheet_names:
            df = pd.read_excel(path, sheet_name=sheet)

            output.append(
                f"\n--- Sheet: {sheet} ---\n"
                f"Rows: {len(df)}\n"
                f"Columns: {list(df.columns)}\n\n"
                f"{df.to_string()}"
            )

        return "\n".join(output)

    except Exception as e:
        return f"[Could not read Excel file: {e}]"


def load_json(path):
    try:
        with open(path, "r", encoding="utf-8") as f:
            data = json.load(f)

        return json.dumps(
            data,
            indent=2,
            ensure_ascii=False
        )

    except Exception as e:
        return f"[Could not read JSON: {e}]"


def load_file(path):
    """
    Determine the file type and extract useful content.
    """

    path = Path(path)
    extension = path.suffix.lower()

    if extension in TEXT_EXTENSIONS:
        return load_text_file(path)

    if extension == ".pdf":
        return load_pdf(path)

    if extension == ".csv":
        return load_csv(path)

    if extension in {".xlsx", ".xls"}:
        return load_excel(path)

    if extension == ".json":
        return load_json(path)

    if extension == ".docx":
        return load_docx(path)

    if extension == ".pptx":
        return load_pptx(path)

    return None

